#!/usr/bin/env python3
"""
RESTfulEzreal - Challenger DuckDB Analytics & Excel Exporter
============================================================
Connects to the RESTfulEzreal DuckDB database (`./output/db/ezreal_data.duckdb`),
combines relational tables (`players`, `match_participants`, `matches`, `match_timelines`)
with deep JSON telemetry (`matches.raw_json` and `match_timelines.raw_json`), and produces:
  1. Console analytical tables
  2. A formatted multi-sheet Excel workbook (`./output/challenger_analysis.xlsx`)
  3. CSV exports (`./output/exports/*.csv`)

Usage:
  python3 scripts/analyze_challenger_db.py [--db ./output/db/ezreal_data.duckdb] [--out ./output/challenger_analysis.xlsx]
"""

import argparse
import json
import math
import os
import subprocess
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd


def query_duckdb(db_path: str, sql: str) -> List[Dict[str, Any]]:
    """Query DuckDB via the python `duckdb` package if installed, else via DuckDB CLI `-json`."""
    try:
        import duckdb  # type: ignore
        conn = duckdb.connect(db_path, read_only=True)
        try:
            df = conn.execute(sql).fetchdf()
            return df.to_dict(orient="records")
        finally:
            conn.close()
    except ImportError:
        pass

    duckdb_bin = "/usr/local/bin/duckdb" if os.path.exists("/usr/local/bin/duckdb") else "duckdb"
    proc = subprocess.run(
        [duckdb_bin, "-readonly", "-json", db_path, sql],
        capture_output=True,
        text=True,
        check=False,
    )
    if proc.returncode != 0 and "lock" in proc.stderr.lower():
        import shutil
        import tempfile
        with tempfile.TemporaryDirectory() as tmpdir:
            snap_db = os.path.join(tmpdir, "snapshot.duckdb")
            shutil.copy2(db_path, snap_db)
            wal_path = db_path + ".wal"
            if os.path.exists(wal_path):
                shutil.copy2(wal_path, snap_db + ".wal")
            proc = subprocess.run(
                [duckdb_bin, "-json", snap_db, sql],
                capture_output=True,
                text=True,
                check=False,
            )
    if proc.returncode != 0:
        raise RuntimeError(f"DuckDB query failed: {proc.stderr.strip()}")
    out = proc.stdout.strip()
    if not out:
        return []
    return json.loads(out)


def safe_div(num: float, den: float, default: float = 0.0) -> float:
    return num / den if den and den != 0 else default


def normalize_role(pos: str) -> str:
    pos_u = (pos or "").strip().upper()
    if pos_u in {"TOP", "JUNGLE", "MIDDLE", "BOTTOM", "UTILITY"}:
        return pos_u
    if pos_u == "MID":
        return "MIDDLE"
    if pos_u in {"BOT", "ADC"}:
        return "BOTTOM"
    if pos_u in {"SUP", "SUPPORT"}:
        return "UTILITY"
    return "UNKNOWN"


def parse_timeline_metrics(raw_tl_str: str) -> Dict[str, Any]:
    """
    Extract per-participant @10m and @15m Gold/XP/CS snapshots and match-level
    first objective timestamps from a Match-v5 Timeline JSON payload.
    """
    result: Dict[str, Any] = {
        "frames_10m": {},
        "frames_15m": {},
        "first_blood_min": None,
        "first_blood_team": None,
        "first_dragon_min": None,
        "first_herald_min": None,
        "first_tower_min": None,
    }
    if not raw_tl_str:
        return result

    try:
        doc = json.loads(raw_tl_str)
    except Exception:
        return result

    info = doc.get("info") or {}
    frames = info.get("frames") or []
    if not frames:
        return result

    def extract_pframes(frame_obj: Dict[str, Any]) -> Dict[int, Dict[str, int]]:
        pframes = frame_obj.get("participantFrames") or {}
        out_map: Dict[int, Dict[str, int]] = {}
        for k, v in pframes.items():
            try:
                pid = int(k)
            except ValueError:
                continue
            minions = int(v.get("minionsKilled", 0)) + int(v.get("jungleMinionsKilled", 0))
            out_map[pid] = {
                "gold": int(v.get("totalGold", 0)),
                "xp": int(v.get("xp", 0)),
                "cs": minions,
                "level": int(v.get("level", 1)),
            }
        return out_map

    if len(frames) > 10:
        result["frames_10m"] = extract_pframes(frames[10])
    if len(frames) > 15:
        result["frames_15m"] = extract_pframes(frames[15])

    # Scan events for First Blood, First Dragon, First Herald, First Tower timestamps
    for frame in frames:
        for ev in frame.get("events") or []:
            etype = ev.get("type")
            ts_min = round(float(ev.get("timestamp", 0)) / 60000.0, 2)
            if etype == "CHAMPION_KILL" and result["first_blood_min"] is None:
                killer_id = int(ev.get("killerId", 0))
                if killer_id > 0:
                    result["first_blood_min"] = ts_min
                    result["first_blood_team"] = 100 if killer_id <= 5 else 200
            elif etype == "ELITE_MONSTER_KILL":
                mtype = ev.get("monsterType")
                if mtype == "DRAGON" and result["first_dragon_min"] is None:
                    result["first_dragon_min"] = ts_min
                elif mtype == "RIFTHERALD" and result["first_herald_min"] is None:
                    result["first_herald_min"] = ts_min
            elif etype == "BUILDING_KILL":
                btype = ev.get("buildingType")
                if btype == "TOWER_BUILDING" and result["first_tower_min"] is None:
                    result["first_tower_min"] = ts_min

    return result


def build_datasets(db_path: str) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[str, int]]:
    """
    Load DuckDB tables and parse Match-v5 + Timeline JSON payloads into:
      - participant_df: 1 row per (match_id, participant_id) with combat, economy, and @10m/@15m lane diffs
      - team_df: 1 row per (match_id, team_id) with objective control and game tempo metrics
      - table_counts: row counts in DuckDB
    """
    counts_rows = query_duckdb(
        db_path,
        """
        SELECT
            (SELECT COUNT(*) FROM players) AS players,
            (SELECT COUNT(*) FROM player_aliases) AS aliases,
            (SELECT COUNT(*) FROM matches) AS matches,
            (SELECT COUNT(*) FROM match_participants) AS match_participants,
            (SELECT COUNT(*) FROM match_timelines) AS match_timelines
        """,
    )
    table_counts = counts_rows[0] if counts_rows else {}

    matches_rows = query_duckdb(
        db_path,
        """
        SELECT
            m.match_id,
            m.game_creation,
            m.game_duration,
            m.game_version,
            m.queue_id,
            CAST(m.raw_json AS VARCHAR) AS match_json,
            CAST(t.raw_json AS VARCHAR) AS timeline_json
        FROM matches m
        LEFT JOIN match_timelines t ON m.match_id = t.match_id
        ORDER BY m.game_creation DESC
        """,
    )

    mp_rows = query_duckdb(
        db_path,
        """
        SELECT
            match_id,
            participant_id,
            player_id,
            game_name,
            tag_line
        FROM match_participants
        """,
    )
    mp_lookup: Dict[Tuple[str, int], Dict[str, Any]] = {}
    for r in mp_rows:
        mp_lookup[(str(r["match_id"]), int(r["participant_id"]))] = r

    participant_records: List[Dict[str, Any]] = []
    team_records: List[Dict[str, Any]] = []

    for mrow in matches_rows:
        match_id = str(mrow["match_id"])
        duration_sec = int(mrow.get("game_duration") or 0)
        duration_min = max(duration_sec / 60.0, 1.0)
        game_version = str(mrow.get("game_version") or "")
        queue_id = int(mrow.get("queue_id") or 0)

        tl_data = parse_timeline_metrics(mrow.get("timeline_json") or "")
        frames_10m = tl_data["frames_10m"]
        frames_15m = tl_data["frames_15m"]
        has_timeline = bool(frames_10m or frames_15m)

        raw_match_str = mrow.get("match_json") or ""
        if not raw_match_str:
            continue
        try:
            mdoc = json.loads(raw_match_str)
        except Exception:
            continue

        info = mdoc.get("info") or {}
        participants = info.get("participants") or []
        teams = info.get("teams") or []

        # Team totals for share computation
        team_kills: Dict[int, int] = {100: 0, 200: 0}
        team_dmg: Dict[int, int] = {100: 0, 200: 0}
        team_gold: Dict[int, int] = {100: 0, 200: 0}

        for p in participants:
            tid = int(p.get("teamId", 100))
            team_kills[tid] = team_kills.get(tid, 0) + int(p.get("kills", 0))
            team_dmg[tid] = team_dmg.get(tid, 0) + int(p.get("totalDamageDealtToChampions", 0))
            team_gold[tid] = team_gold.get(tid, 0) + int(p.get("goldEarned", 0))

        # Parse team-level objective stats
        for t in teams:
            tid = int(t.get("teamId", 100))
            objs = t.get("objectives") or {}
            champ_obj = objs.get("champion") or {}
            drag_obj = objs.get("dragon") or {}
            herald_obj = objs.get("riftHerald") or {}
            horde_obj = objs.get("horde") or {}
            tower_obj = objs.get("tower") or {}
            baron_obj = objs.get("baron") or {}

            team_records.append(
                {
                    "match_id": match_id,
                    "queue_id": queue_id,
                    "game_version": game_version,
                    "duration_min": round(duration_min, 2),
                    "team_id": tid,
                    "side": "Blue (100)" if tid == 100 else "Red (200)",
                    "win": 1 if bool(t.get("win", False)) else 0,
                    "team_kills": team_kills.get(tid, 0),
                    "team_gold": team_gold.get(tid, 0),
                    "first_blood": 1 if bool(champ_obj.get("first", False)) else 0,
                    "first_dragon": 1 if bool(drag_obj.get("first", False)) else 0,
                    "dragon_kills": int(drag_obj.get("kills", 0)),
                    "first_herald": 1 if bool(herald_obj.get("first", False)) else 0,
                    "herald_kills": int(herald_obj.get("kills", 0)),
                    "voidgrub_kills": int(horde_obj.get("kills", 0)),
                    "first_tower": 1 if bool(tower_obj.get("first", False)) else 0,
                    "tower_kills": int(tower_obj.get("kills", 0)),
                    "first_baron": 1 if bool(baron_obj.get("first", False)) else 0,
                    "baron_kills": int(baron_obj.get("kills", 0)),
                    "first_blood_min": tl_data["first_blood_min"],
                    "first_dragon_min": tl_data["first_dragon_min"],
                    "first_tower_min": tl_data["first_tower_min"],
                }
            )

        # Map (team_id, role) -> participant index for lane opponent pairing
        role_map: Dict[Tuple[int, str], Dict[str, Any]] = {}
        parsed_parts: List[Dict[str, Any]] = []

        for idx, p in enumerate(participants):
            pid = int(p.get("participantId") or (idx + 1))
            tid = int(p.get("teamId", 100 if pid <= 5 else 200))
            role = normalize_role(p.get("teamPosition") or p.get("individualPosition") or "")

            mp_info = mp_lookup.get((match_id, pid), {})
            gname = str(p.get("riotIdGameName") or mp_info.get("game_name") or p.get("summonerName") or f"Player_{pid}")
            tline = str(p.get("riotIdTagline") or mp_info.get("tag_line") or "")
            riot_id = f"{gname}#{tline}" if tline else gname
            player_id = str(mp_info.get("player_id") or "")

            kills = int(p.get("kills", 0))
            deaths = int(p.get("deaths", 0))
            assists = int(p.get("assists", 0))
            kda = (kills + assists) / max(deaths, 1)
            dmg = int(p.get("totalDamageDealtToChampions", 0))
            gold = int(p.get("goldEarned", 0))
            cs = int(p.get("totalMinionsKilled", 0)) + int(p.get("neutralMinionsKilled", 0))
            vision = int(p.get("visionScore", 0))
            pink_wards = int(p.get("visionWardsBoughtInGame", 0))
            turret_plates = int((p.get("challenges") or {}).get("turretPlatesTaken", 0))

            s10 = frames_10m.get(pid, {})
            s15 = frames_15m.get(pid, {})

            rec: Dict[str, Any] = {
                "match_id": match_id,
                "game_version": game_version,
                "queue_id": queue_id,
                "duration_min": round(duration_min, 2),
                "participant_id": pid,
                "player_id": player_id,
                "riot_id": riot_id,
                "game_name": gname,
                "tag_line": tline,
                "team_id": tid,
                "role": role,
                "champion": str(p.get("championName") or "Unknown"),
                "champion_id": int(p.get("championId") or 0),
                "win": 1 if bool(p.get("win", False)) else 0,
                "kills": kills,
                "deaths": deaths,
                "assists": assists,
                "kda": round(kda, 2),
                "damage_to_champions": dmg,
                "dpm": round(dmg / duration_min, 1),
                "dmg_share_pct": round(100.0 * safe_div(dmg, team_dmg.get(tid, 0)), 1),
                "gold_earned": gold,
                "gpm": round(gold / duration_min, 1),
                "gold_share_pct": round(100.0 * safe_div(gold, team_gold.get(tid, 0)), 1),
                "cs": cs,
                "cs_per_min": round(cs / duration_min, 2),
                "vision_score": vision,
                "vision_per_min": round(vision / duration_min, 2),
                "control_wards": pink_wards,
                "turret_plates": turret_plates,
                "kp_pct": round(100.0 * safe_div(kills + assists, team_kills.get(tid, 0)), 1),
                "has_timeline": has_timeline,
                "gold_10m": s10.get("gold"),
                "xp_10m": s10.get("xp"),
                "cs_10m": s10.get("cs"),
                "gold_15m": s15.get("gold"),
                "xp_15m": s15.get("xp"),
                "cs_15m": s15.get("cs"),
                "gold_diff_10m": None,
                "xp_diff_10m": None,
                "cs_diff_10m": None,
                "gold_diff_15m": None,
                "xp_diff_15m": None,
                "cs_diff_15m": None,
            }
            if role != "UNKNOWN":
                role_map[(tid, role)] = rec
            parsed_parts.append(rec)

        # Compute direct lane opponent differentials (@10m and @15m)
        for rec in parsed_parts:
            role = rec["role"]
            tid = rec["team_id"]
            opp_tid = 200 if tid == 100 else 100
            opp = role_map.get((opp_tid, role))
            if opp:
                if rec["gold_10m"] is not None and opp["gold_10m"] is not None:
                    rec["gold_diff_10m"] = rec["gold_10m"] - opp["gold_10m"]
                    rec["xp_diff_10m"] = rec["xp_10m"] - opp["xp_10m"]
                    rec["cs_diff_10m"] = rec["cs_10m"] - opp["cs_10m"]
                if rec["gold_15m"] is not None and opp["gold_15m"] is not None:
                    rec["gold_diff_15m"] = rec["gold_15m"] - opp["gold_15m"]
                    rec["xp_diff_15m"] = rec["xp_15m"] - opp["xp_15m"]
                    rec["cs_diff_15m"] = rec["cs_15m"] - opp["cs_15m"]

        participant_records.extend(parsed_parts)

    return pd.DataFrame(participant_records), pd.DataFrame(team_records), table_counts


def compute_analysis_tables(
    part_df: pd.DataFrame,
    team_df: pd.DataFrame,
    table_counts: Dict[str, int],
) -> Dict[str, pd.DataFrame]:
    """Compute the 6 analytical DataFrames for console and Excel export."""
    if part_df.empty:
        return {}

    # 1. Overview KPIs
    total_matches = int(part_df["match_id"].nunique())
    ranked_solo_matches = int(part_df[part_df["queue_id"] == 420]["match_id"].nunique())
    tl_matches = int(part_df[part_df["has_timeline"] == True]["match_id"].nunique())
    avg_duration = round(float(part_df.groupby("match_id")["duration_min"].first().mean()), 2)

    blue_wr = 0.0
    fb_wr = 0.0
    fdrag_wr = 0.0
    ftower_wr = 0.0
    fherald_wr = 0.0
    fbaron_wr = 0.0
    if not team_df.empty:
        blue_df = team_df[team_df["team_id"] == 100]
        if not blue_df.empty:
            blue_wr = round(100.0 * float(blue_df["win"].mean()), 1)
        for col, attr in [
            ("first_blood", "fb"),
            ("first_dragon", "fdrag"),
            ("first_tower", "ftower"),
            ("first_herald", "fherald"),
            ("first_baron", "fbaron"),
        ]:
            sub = team_df[team_df[col] == 1]
            val = round(100.0 * float(sub["win"].mean()), 1) if not sub.empty else 0.0
            if attr == "fb":
                fb_wr = val
            elif attr == "fdrag":
                fdrag_wr = val
            elif attr == "ftower":
                ftower_wr = val
            elif attr == "fherald":
                fherald_wr = val
            elif attr == "fbaron":
                fbaron_wr = val

    overview_df = pd.DataFrame(
        [
            {"Metric": "Canonical Players in DuckDB (`players`)", "Value": table_counts.get("players", 0)},
            {"Metric": "Player Riot ID Aliases (`player_aliases`)", "Value": table_counts.get("aliases", 0)},
            {"Metric": "Total Matches in DuckDB (`matches`)", "Value": total_matches},
            {"Metric": "Ranked Solo 5x5 Matches (`queue_id = 420`)", "Value": ranked_solo_matches},
            {"Metric": "Match Timelines in DuckDB (`match_timelines`)", "Value": table_counts.get("match_timelines", 0)},
            {"Metric": "Matches with @10m/@15m Timeline Frames", "Value": tl_matches},
            {"Metric": "Total Participant Performance Rows", "Value": len(part_df)},
            {"Metric": "Average Match Duration (Minutes)", "Value": avg_duration},
            {"Metric": "Blue Side Win Rate (%)", "Value": blue_wr},
            {"Metric": "First Blood Win Conversion Rate (%)", "Value": fb_wr},
            {"Metric": "First Dragon Win Conversion Rate (%)", "Value": fdrag_wr},
            {"Metric": "First Herald Win Conversion Rate (%)", "Value": fherald_wr},
            {"Metric": "First Tower Win Conversion Rate (%)", "Value": ftower_wr},
            {"Metric": "First Baron Win Conversion Rate (%)", "Value": fbaron_wr},
        ]
    )

    # 2. Role Impact & Early-Game Snowball (@10m & @15m)
    role_order = ["TOP", "JUNGLE", "MIDDLE", "BOTTOM", "UTILITY"]
    role_rows: List[Dict[str, Any]] = []
    for role in role_order:
        rdf = part_df[part_df["role"] == role]
        if rdf.empty:
            continue
        tl_rdf = rdf.dropna(subset=["gold_diff_15m"])
        ahead_15 = tl_rdf[tl_rdf["gold_diff_15m"] > 0]
        behind_15 = tl_rdf[tl_rdf["gold_diff_15m"] < 0]
        big_lead_15 = tl_rdf[tl_rdf["gold_diff_15m"] >= 500]

        tl_10_rdf = rdf.dropna(subset=["gold_diff_10m"])
        ahead_10 = tl_10_rdf[tl_10_rdf["gold_diff_10m"] > 0]

        role_rows.append(
            {
                "Role": role,
                "Appearances": len(rdf),
                "Timeline_Sample": len(tl_rdf),
                "Avg_KDA": round(float(rdf["kda"].mean()), 2),
                "Avg_DPM": round(float(rdf["dpm"].mean()), 1),
                "Avg_Dmg_Share_%": round(float(rdf["dmg_share_pct"].mean()), 1),
                "Avg_GPM": round(float(rdf["gpm"].mean()), 1),
                "Avg_Gold_Share_%": round(float(rdf["gold_share_pct"].mean()), 1),
                "Avg_CS_Per_Min": round(float(rdf["cs_per_min"].mean()), 2),
                "Avg_Vision_Per_Min": round(float(rdf["vision_per_min"].mean()), 2),
                "Avg_KP_%": round(float(rdf["kp_pct"].mean()), 1),
                "Avg_Turret_Plates": round(float(rdf["turret_plates"].mean()), 2),
                "Avg_Abs_Gold_Diff_10m": round(float(tl_10_rdf["gold_diff_10m"].abs().mean()), 0) if not tl_10_rdf.empty else None,
                "Win_Rate_Ahead_10m_%": round(100.0 * float(ahead_10["win"].mean()), 1) if not ahead_10.empty else None,
                "Avg_Abs_Gold_Diff_15m": round(float(tl_rdf["gold_diff_15m"].abs().mean()), 0) if not tl_rdf.empty else None,
                "Win_Rate_Ahead_15m_%": round(100.0 * float(ahead_15["win"].mean()), 1) if not ahead_15.empty else None,
                "Win_Rate_Behind_15m_Comeback_%": round(100.0 * float(behind_15["win"].mean()), 1) if not behind_15.empty else None,
                "Win_Rate_+500g_Lead_15m_%": round(100.0 * float(big_lead_15["win"].mean()), 1) if not big_lead_15.empty else None,
            }
        )
    role_df = pd.DataFrame(role_rows)

    # 3. Champion Meta Stats
    champ_rows: List[Dict[str, Any]] = []
    for champ, cdf in part_df.groupby("champion"):
        games = len(cdf)
        primary_role = cdf["role"].mode().iloc[0] if not cdf["role"].mode().empty else "UNKNOWN"
        tl_cdf = cdf.dropna(subset=["gold_diff_15m"])
        champ_rows.append(
            {
                "Champion": champ,
                "Primary_Role": primary_role,
                "Games": games,
                "Pick_Rate_%": round(100.0 * games / max(total_matches, 1), 1),
                "Win_Rate_%": round(100.0 * float(cdf["win"].mean()), 1),
                "Avg_Kills": round(float(cdf["kills"].mean()), 1),
                "Avg_Deaths": round(float(cdf["deaths"].mean()), 1),
                "Avg_Assists": round(float(cdf["assists"].mean()), 1),
                "Avg_KDA": round(float(cdf["kda"].mean()), 2),
                "Avg_DPM": round(float(cdf["dpm"].mean()), 1),
                "Avg_Dmg_Share_%": round(float(cdf["dmg_share_pct"].mean()), 1),
                "Avg_GPM": round(float(cdf["gpm"].mean()), 1),
                "Avg_CS_Per_Min": round(float(cdf["cs_per_min"].mean()), 2),
                "Avg_KP_%": round(float(cdf["kp_pct"].mean()), 1),
                "Avg_Gold_Diff_15m": round(float(tl_cdf["gold_diff_15m"].mean()), 0) if not tl_cdf.empty else None,
                "Avg_XP_Diff_15m": round(float(tl_cdf["xp_diff_15m"].mean()), 0) if not tl_cdf.empty else None,
            }
        )
    champ_df = pd.DataFrame(champ_rows).sort_values(by=["Games", "Win_Rate_%"], ascending=[False, False]).reset_index(drop=True)

    # 4. Challenger Player Profiles
    player_rows: List[Dict[str, Any]] = []
    for riot_id, pdf in part_df.groupby("riot_id"):
        if not riot_id or riot_id.startswith("Player_"):
            continue
        games = len(pdf)
        non_empty_pids = [p for p in pdf["player_id"].unique() if p]
        canonical_pid = non_empty_pids[0] if non_empty_pids else ""
        primary_role = pdf["role"].mode().iloc[0] if not pdf["role"].mode().empty else "UNKNOWN"
        top_champs = ", ".join(
            f"{c} ({cnt})" for c, cnt in pdf["champion"].value_counts().head(3).items()
        )
        tl_pdf = pdf.dropna(subset=["gold_diff_15m"])
        player_rows.append(
            {
                "Riot_ID": riot_id,
                "Canonical_Player_ID": canonical_pid,
                "Games": games,
                "Win_Rate_%": round(100.0 * float(pdf["win"].mean()), 1),
                "Avg_KDA": round(float(pdf["kda"].mean()), 2),
                "Avg_Kills": round(float(pdf["kills"].mean()), 1),
                "Avg_Deaths": round(float(pdf["deaths"].mean()), 1),
                "Avg_Assists": round(float(pdf["assists"].mean()), 1),
                "Avg_DPM": round(float(pdf["dpm"].mean()), 1),
                "Avg_Dmg_Share_%": round(float(pdf["dmg_share_pct"].mean()), 1),
                "Avg_GPM": round(float(pdf["gpm"].mean()), 1),
                "Avg_CS_Per_Min": round(float(pdf["cs_per_min"].mean()), 2),
                "Avg_Vision_Per_Min": round(float(pdf["vision_per_min"].mean()), 2),
                "Avg_KP_%": round(float(pdf["kp_pct"].mean()), 1),
                "Avg_Gold_Diff_15m": round(float(tl_pdf["gold_diff_15m"].mean()), 0) if not tl_pdf.empty else None,
                "Primary_Role": primary_role,
                "Top_Champions": top_champs,
            }
        )
    player_df = pd.DataFrame(player_rows).sort_values(
        by=["Games", "Win_Rate_%", "Avg_KDA"], ascending=[False, False, False]
    ).reset_index(drop=True)

    # 5. Objective Tempo & Win Conversion Table
    obj_rows: List[Dict[str, Any]] = []
    if not team_df.empty:
        conditions = [
            ("First Blood Secured", team_df["first_blood"] == 1),
            ("First Dragon Secured", team_df["first_dragon"] == 1),
            ("First Rift Herald Secured", team_df["first_herald"] == 1),
            ("First Tower Secured", team_df["first_tower"] == 1),
            ("3+ Voidgrubs Secured", team_df["voidgrub_kills"] >= 3),
            ("5+ Voidgrubs Secured", team_df["voidgrub_kills"] >= 5),
            ("First Baron Secured", team_df["first_baron"] == 1),
            ("3+ Dragons Secured", team_df["dragon_kills"] >= 3),
            ("4+ Dragons (Soul Point / Soul)", team_df["dragon_kills"] >= 4),
            ("First Blood + First Tower + First Dragon", (team_df["first_blood"] == 1) & (team_df["first_tower"] == 1) & (team_df["first_dragon"] == 1)),
        ]
        for label, mask in conditions:
            sub = team_df[mask]
            if sub.empty:
                continue
            obj_rows.append(
                {
                    "Objective_Condition": label,
                    "Team_Instances": len(sub),
                    "Frequency_%": round(100.0 * len(sub) / len(team_df), 1),
                    "Win_Rate_%": round(100.0 * float(sub["win"].mean()), 1),
                    "Avg_Game_Duration_Min": round(float(sub["duration_min"].mean()), 2),
                    "Avg_Team_Kills": round(float(sub["team_kills"].mean()), 1),
                    "Avg_Towers_Taken": round(float(sub["tower_kills"].mean()), 1),
                }
            )
    obj_df = pd.DataFrame(obj_rows)

    return {
        "Overview_KPIs": overview_df,
        "Role_Snowball_10m_15m": role_df,
        "Objective_Tempo_Impact": obj_df,
        "Champion_Meta_Stats": champ_df,
        "Challenger_Player_Profiles": player_df,
        "Participant_Match_Log": part_df,
    }


def export_to_excel_and_csv(tables: Dict[str, pd.DataFrame], xlsx_path: Path, csv_dir: Path) -> None:
    """Write formatted multi-sheet Excel workbook via openpyxl and CSV files."""
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    csv_dir.mkdir(parents=True, exist_ok=True)

    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill(start_color="1A181D", end_color="1A181D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="EBAE03")  # Gold accent matching UI
    body_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        for sheet_name, df in tables.items():
            # Save CSV copy
            df.to_csv(csv_dir / f"{sheet_name.lower()}.csv", index=False)

            # Write Excel sheet
            df.to_excel(writer, sheet_name=sheet_name[:31], index=False)
            ws = writer.sheets[sheet_name[:31]]
            ws.views.sheetView[0].showGridLines = True
            ws.freeze_panes = "A2"

            # Format header row
            for col_idx in range(1, len(df.columns) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Format data cells & auto-fit column widths
            max_rows_to_style = min(len(df) + 1, 2000)
            for col_idx, col_name in enumerate(df.columns, start=1):
                col_letter = get_column_letter(col_idx)
                max_len = len(str(col_name))
                for row_idx in range(2, max_rows_to_style + 1):
                    c = ws.cell(row=row_idx, column=col_idx)
                    c.font = body_font
                    c.border = thin_border
                    val_str = str(c.value) if c.value is not None else ""
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 42)


def main() -> int:
    parser = argparse.ArgumentParser(description="Analyze RESTfulEzreal Challenger DuckDB database and export to Excel/CSV.")
    parser.add_argument("--db", default="./output/db/ezreal_data.duckdb", help="Path to DuckDB database file")
    parser.add_argument("--out", default="./output/challenger_analysis.xlsx", help="Output Excel (.xlsx) path")
    parser.add_argument("--csv-dir", default="./output/exports", help="Directory for CSV exports")
    args = parser.parse_args()

    if not os.path.exists(args.db):
        print(f"[ERROR] DuckDB database not found at: {args.db}", file=sys.stderr)
        return 1

    print(f"[Analytics] Loading DuckDB database: {args.db}")
    part_df, team_df, table_counts = build_datasets(args.db)
    if part_df.empty:
        print("[WARN] No match data found in DuckDB.")
        return 0

    tables = compute_analysis_tables(part_df, team_df, table_counts)

    print("\n=== 1. DATASET & OBJECTIVE OVERVIEW ===")
    print(tables["Overview_KPIs"].to_string(index=False))

    print("\n=== 2. ROLE IMPACT & EARLY-GAME SNOWBALL (@10M / @15M) ===")
    print(tables["Role_Snowball_10m_15m"].to_string(index=False))

    print("\n=== 3. OBJECTIVE TEMPO & WIN CONVERSION ===")
    print(tables["Objective_Tempo_Impact"].to_string(index=False))

    print("\n=== 4. TOP 15 MOST PLAYED CHAMPIONS IN DATASET ===")
    print(tables["Champion_Meta_Stats"].head(15).to_string(index=False))

    print("\n=== 5. TOP 15 MOST ACTIVE PLAYERS IN DATASET ===")
    print(tables["Challenger_Player_Profiles"].head(15).to_string(index=False))

    xlsx_path = Path(args.out)
    csv_dir = Path(args.csv_dir)
    export_to_excel_and_csv(tables, xlsx_path, csv_dir)
    print(f"\n[Analytics] Exported multi-sheet Excel workbook to: {xlsx_path}")
    print(f"[Analytics] Exported 6 CSV tables to: {csv_dir}/")
    return 0


if __name__ == "__main__":
    sys.exit(main())
